"""
tests/test_offline.py — bina API key poora system test (offline demo mode)
=========================================================================

Chalao:  python tests/test_offline.py       (standalone acceptance; pytest -q runs the unit suite)

Ye test real cheezein verify karta hai:
  ✅ memory store (SQLite + FTS5 recall)
  ✅ notes/notepad save + index
  ✅ code engine — asli code likhta/ chalata hai (offline templates) + artifacts banate hain
  ✅ browser tools — offline plan + mocked search/fetch fallback
  ✅ DSPy modules — route/plan/extract fallbacks
  ✅ LangGraph supervisor — poora graph end-to-end chalta hai, workers execute hote hain
  ✅ identity migration — koi purana naam bacha nahi
"""
from __future__ import annotations

import asyncio
import json
import os
import socket
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# Filled only after main() has redirected every configured data path to TemporaryDirectory.
WORKSPACE_DIR: Path | None = None
NOTES_DIR: Path | None = None
MEMORY_DB: Path | None = None
FIXTURES_DIR: Path | None = None
_WEB_SEARCH_CALLS: list[tuple[str, int]] = []
_WEB_FETCH_CALLS: list[tuple[str, int]] = []


def _configure_selftest_environment(temp_root: Path) -> dict[str, str | None]:
    workspace = temp_root / "workspace"
    values = {
        "BRAIN_PROVIDER": "demo",
        "OFFLINE_MODE": "1",
        "MODEL_DISCOVERY": "0",
        "WORKSPACE_DIR": str(workspace),
        "NOTES_DIR": str(workspace / "notes"),
        "MEMORY_DB": str(workspace / "nikhilai_memory.sqlite3"),
        "SESSION_STATE_FILE": str(workspace / "nikhilai_memory.json"),
        "RUNTIME_SETTINGS": str(workspace / "runtime_settings.json"),
        "AUTOMATION_FILE": str(workspace / "automations.json"),
        "APP_MEMORY_FILE": str(workspace / "app_launcher_memory.json"),
        "FILE_SEARCH_ROOTS": str(workspace),
        "BROWSER_PROFILE_DIR": str(workspace / "browser-profile"),
        "JARVIS_EXEC_APPROVAL": "risky",
    }
    tracked = {*values, "JARVIS_SECRET_demo"}
    previous = {key: os.environ.get(key) for key in tracked}
    os.environ.update(values)
    return previous


def _restore_environment(previous: dict[str, str | None]) -> None:
    for key, value in previous.items():
        if value is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = value


def _install_network_guard() -> tuple[list[str], object]:
    """Block and record all socket/DNS attempts for the complete standalone selftest."""
    attempts: list[str] = []
    originals: list[tuple[object, str, object]] = []

    def denied(*_args, **_kwargs):
        attempts.append("socket/DNS")
        raise OSError("live network is disabled in the offline selftest")

    for owner, name in (
        (socket, "create_connection"),
        (socket, "getaddrinfo"),
        (socket.socket, "connect"),
        (socket.socket, "connect_ex"),
        (socket.socket, "sendto"),
    ):
        if hasattr(owner, name):
            originals.append((owner, name, getattr(owner, name)))
            setattr(owner, name, denied)

    def restore() -> None:
        for owner, name, value in reversed(originals):
            setattr(owner, name, value)

    return attempts, restore


def _is_within(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def _temporary_parent() -> Path:
    """Pick a system temp directory outside the checkout before creating any test data."""
    candidates = (Path(tempfile.gettempdir()), Path("/tmp"), Path("/var/tmp"), Path.home())
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
            if resolved.is_dir() and not _is_within(resolved, ROOT):
                return resolved
        except OSError:
            continue
    raise RuntimeError("No temporary directory outside the repository is available")


async def _mock_web_search(query: str, num: int = 5) -> list[dict[str, str]]:
    _WEB_SEARCH_CALLS.append((query, num))
    return [{"title": "Mock offline source", "url": "https://example.invalid/selftest-page",
             "snippet": "Mock search result; no provider was contacted."}]


async def _mock_web_fetch(url: str, max_chars: int = 12000, timeout: float = 25.0) -> str:
    _WEB_FETCH_CALLS.append((url, max_chars))
    return "MOCKED_OFFLINE_FETCH_CONTENT — no network service was contacted."


def _mock_weather(_city: str = "") -> str:
    return "Mock offline weather for Kota; no weather service was contacted."

PASS, FAIL = "✅", "❌"
_results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    _results.append((name, ok, detail))
    print(f"  {PASS if ok else FAIL} {name}" + (f" — {detail}" if detail else ""))


async def t_memory() -> None:
    from jarvis.memory.store import get_memory

    mem = get_memory()
    fid = mem.remember_fact("Nikhil Mehra ko filter coffee pasand hai", key="coffee", tags="owner,test")
    hits = mem.search("coffee")
    check("memory: fact save + FTS recall", bool(fid) and any("coffee" in h.text for h in hits),
          f"{len(hits)} hits")
    mem.add_turn("test user bolo", "test assistant jawab", session="selftest")
    ctx = mem.context_block(limit=4, session="selftest")
    check("memory: conversation context", "test user bolo" in ctx, f"{len(ctx)} chars")
    stats = mem.stats()
    check("memory: stats", stats["facts"] > 0 and stats["episodes"] > 0, json.dumps(stats))


async def t_notes() -> None:
    from jarvis.tools.file_tools import invalidate_index
    from jarvis.tools.notes_tools import _note_path, list_notes, read_note, save_note

    name = f"selftest_note_{time.time_ns()}"
    try:
        out = await save_note.ainvoke({"name": name, "content": "Jarvis offline selftest note", "open_in_notepad": False})
        check("notes: save_note", "save ho gaya" in out or "saved" in out.lower(), out.splitlines()[0][:80])
        read = await read_note.ainvoke({"name": name})
        check("notes: read_note", "selftest note" in read, read.splitlines()[0][:80])
        listing = await list_notes.ainvoke({})
        check("notes: list_notes", name in listing, f"{len(listing.splitlines())} lines")
    finally:
        _note_path(name).unlink(missing_ok=True)
        invalidate_index()


async def t_code_engine() -> None:
    from jarvis.tools.code_engine import code_engine_run, run_python, static_check

    ok_safe, _ = static_check("import os\nos.system('rm -rf /')", "python")
    check("code: dangerous code blocked", not ok_safe, "rm -rf / blocked")

    ok_syntax, msg = static_check("def f(:\n  pass", "python")
    check("code: syntax error detected", not ok_syntax, msg[:60])

    res = await run_python.ainvoke({"code": "print('jarvis', 6*7)"})
    check("code: run_python executes", "42" in res, res.strip().splitlines()[-1][:60])

    outcome = await code_engine_run(task="ek csv report bana do sales data ka", language="python", run=True)
    check("code: engine draft+run+heal", outcome["ok"] and Path(outcome["script"]).exists(),
          f"attempts={outcome['attempts']} script={Path(outcome['script']).name}")


async def t_browser() -> None:
    from unittest.mock import patch

    from jarvis.tools import browser_tools

    status = browser_tools.browser_status()
    check("browser: status report", isinstance(status, dict) and "mode" in status, json.dumps(status)[:110])

    doc = await browser_tools.browser_doctor.ainvoke({})
    check("browser: doctor", "browser-use" in doc, doc.splitlines()[1][:70])

    async def no_live_browser_agent(*_args, **_kwargs):
        return None

    search_before, fetch_before = len(_WEB_SEARCH_CALLS), len(_WEB_FETCH_CALLS)
    with patch.object(browser_tools, "build_browser_agent", no_live_browser_agent):
        outcome = await browser_tools.run_browser_task(task="browser se news laao", start_url="")
    result = str(outcome.get("result", ""))
    mocked_search_and_fetch = (
        len(_WEB_SEARCH_CALLS) > search_before
        and len(_WEB_FETCH_CALLS) > fetch_before
        and "MOCKED_OFFLINE_FETCH_CONTENT" in result
    )
    check("browser: fallback uses mocked search + fetch only",
          isinstance(outcome, dict) and "result" in outcome and mocked_search_and_fetch,
          f"search={len(_WEB_SEARCH_CALLS) - search_before} fetch={len(_WEB_FETCH_CALLS) - fetch_before} "
          f"offline={outcome.get('offline')}")


async def t_dspy() -> None:
    from jarvis.dspy_engine import modules as M

    r = M.route("notepad me laptop_prices naam se ye save kardo")
    check("dspy: router fallback", "desktop" in (r.get("workers") or []), str(r.get("workers")))
    r2 = M.route("alag alag website se price laao")
    check("dspy: browser routing", "browser" in (r2.get("workers") or []), str(r2.get("workers")))
    p = M.plan("price compare karo", ["browser", "desktop"])
    check("dspy: planner", len(p.get("steps", [])) >= 3, f"{len(p.get('steps', []))} steps")
    ex = M.extract_structured("Price: ₹54,999\nRating: 4.3", "https://x.com", ["price", "rating"])
    check("dspy: extractor fallback", ex.get("price") is not None, json.dumps(ex)[:110])


async def t_graph() -> None:
    from jarvis.graph.build import get_jarvis_graph, jarvis_ainvoke, jarvis_stream

    graph = get_jarvis_graph()
    check("graph: compiles", graph is not None)

    # 1) notes rakhne wala desktop task; unique filename means reruns never overwrite owner data.
    from jarvis.config import NOTES_DIR
    from jarvis.tools.file_tools import invalidate_index

    note_name = f"mtg_selftest_{time.time_ns()}"
    try:
        r1 = await jarvis_ainvoke(
            f"jarvis, notes me {note_name} naam se ye save kardo: client call 5 baje", session_id="t1"
        )
        answer = str(r1.get("answer") or "")
        requirements_done = bool(r1.get("requirements")) and all(
            item.get("status") == "done" for item in r1.get("requirements", {}).values()
        )
        answer_status_agree = (
            r1.get("status") == "done"
            and not r1.get("requirements_unresolved")
            and requirements_done
            and ("Verified ✅" in answer)
        )
        check("graph: desktop/notes task end-to-end", bool(answer) and answer_status_agree,
              f"status={r1.get('status')} requirements_done={requirements_done} "
              f"answer={answer.splitlines()[0][:90]!r}")
    finally:
        (NOTES_DIR / f"{note_name}.md").unlink(missing_ok=True)
        invalidate_index()

    # 2) code task (offline template)
    r2 = await jarvis_ainvoke("ek script likho jo csv report bana de", session_id="t2")
    check("graph: code worker", any(f.get("worker") == "code" for f in r2.get("findings", [])),
          f"workers={r2.get('workers')}")

    # 3) streaming events
    evs = []
    async for ev in jarvis_stream("mausam batao Kota ka", session_id="t3"):
        evs.append(ev)
        if ev.get("node") == "final":
            break
    check("graph: streaming events", len(evs) > 3 and any(e.get("node") == "final" for e in evs),
          f"{len(evs)} events")

    # 4) identity
    r4 = await jarvis_ainvoke("tum kaun ho?", session_id="t4")
    ans = (r4.get("answer") or "").lower()
    check("graph: identity answer", "jarvis" in ans or "nikhil" in ans, ans[:70].replace("\n", " "))


def t_identity() -> None:
    """Purana naam poore project me kahin nahi hona chahiye (legacy_original ko chhod kar)."""
    # needles concatenated on purpose (taaki ye test file khud ko match na kare)
    bad = ("sach" + "deva", "gau" + "rav", "gaur" + "v")
    offenders: list[str] = []
    for path in ROOT.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".py", ".md", ".txt", ".json", ".toml", ".env", ".example", ".html"}:
            continue
        # legacy_original = migrated reference dump; rename tool & this test contain the
        # old needles by design (they are the migration machinery itself).
        if (
            "legacy_original" in path.parts
            or "rename_report" in path.name
            or path.name in {"rename_identities.py", "test_offline.py"}
        ):
            continue
        try:
            text = path.read_text(encoding="utf-8", errors="ignore").lower()
        except Exception:
            continue
        if any(b in text for b in bad):
            offenders.append(str(path.relative_to(ROOT)))
    check("identity: no old names in new code", not offenders, f"offenders={offenders[:4]}")




# ===========================================================================
#  v2.1 modules: app finder, any-file reader, REPL, session state, processes,
#  approval gate, skills, loop guard
# ===========================================================================
def t_app_finder() -> None:
    from jarvis.tools import app_finder

    fixture_apps = [
        {"name": "Google Chrome", "path": "desktop:google-chrome", "kind": "desktop", "extra": ""},
        {"name": "ImageMagick", "path": "desktop:imagemagick", "kind": "desktop", "extra": "image viewer"},
    ]
    original_cache = dict(app_finder._INDEX_CACHE)
    original_flags = (app_finder.IS_WINDOWS, app_finder.IS_MAC)
    original_indexer = app_finder._index_linux

    def fake_linux_index(apps: list[dict]) -> None:
        for candidate in fixture_apps:
            app_finder._add(apps, candidate["name"], candidate["path"], candidate["kind"], candidate["extra"])

    app_finder.IS_WINDOWS = False
    app_finder.IS_MAC = False
    app_finder._index_linux = fake_linux_index
    app_finder._INDEX_CACHE = {"ts": 0.0, "apps": []}
    try:
        apps = app_finder.build_app_index(force=True)
        matches = app_finder.find_apps("browser")
    finally:
        app_finder.IS_WINDOWS, app_finder.IS_MAC = original_flags
        app_finder._index_linux = original_indexer
        app_finder._INDEX_CACHE = original_cache

    check("app_finder: synthetic index builds", len(apps) == 2, f"{len(apps)} fixture apps indexed")
    check("app_finder: 'browser' resolves to browser candidate",
          bool(matches) and matches[0]["name"] == "Google Chrome"
          and all(match["name"] != "ImageMagick" for match in matches),
          f"'browser' → {[match['name'] for match in matches[:3]] or 'no match'}")


async def t_file_readers() -> None:
    from jarvis.tools.file_readers import detect_type, read_any_file, supported_types

    assert FIXTURES_DIR is not None
    fixtures = FIXTURES_DIR
    fixtures.mkdir(parents=True, exist_ok=True)

    # text
    txt = fixtures / "sample.txt"
    txt.write_text("hello jarvis\nsecond line", encoding="utf-8")
    check("readers: extension/content type detection", detect_type(txt) == "text", detect_type(txt))
    r_txt = read_any_file(txt)
    check("readers: text file", r_txt["ok"] and "second line" in r_txt["text"], f"kind={r_txt['kind']}")

    # csv
    csv_p = fixtures / "sample.csv"
    csv_p.write_text("model,price\nDell,54990\nHP,49990\n", encoding="utf-8")
    r_csv = read_any_file(csv_p)
    check("readers: csv columns", r_csv["ok"] and "model" in r_csv["text"] and "49990" in r_csv["text"],
          f"kind={r_csv['kind']}")

    # json
    json_p = fixtures / "sample.json"
    json_p.write_text('{"name": "Nikhil", "role": "owner"}', encoding="utf-8")
    r_json = read_any_file(json_p)
    check("readers: json", r_json["ok"] and "Nikhil" in r_json["text"], f"kind={r_json['kind']}")

    # pdf — reportlab ho to asli content check, warna graceful-failure check
    pdf_p = fixtures / "sample.pdf"
    real_pdf = True
    try:
        from reportlab.pdfgen import canvas  # type: ignore

        c = canvas.Canvas(str(pdf_p))
        c.drawString(100, 750, "Jarvis PDF reader test — price 54990")
        c.save()
    except Exception:
        real_pdf = False
        pdf_p.write_bytes(b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>")
    r_pdf = read_any_file(pdf_p)
    if real_pdf:
        check("readers: pdf (asli text)", r_pdf["kind"] == "pdf" and r_pdf["ok"] and "54990" in r_pdf["text"],
              f"kind={r_pdf['kind']} len={len(r_pdf['text'])}")
    else:
        # reportlab nahi hai → reader ko crash nahi karna chahiye, saaf message dena chahiye
        check("readers: pdf (graceful jab reportlab nahi)", r_pdf["kind"] == "pdf" and isinstance(r_pdf["text"], str),
              f"kind={r_pdf['kind']} len={len(r_pdf['text'])} ok={r_pdf['ok']}")

    # docx
    docx_p = fixtures / "sample.docx"
    try:
        import docx  # type: ignore

        d = docx.Document()
        d.add_heading("Jarvis DOCX test", level=1)
        d.add_paragraph("This is a test paragraph with GSTIN 22AAAAA0000A1Z5")
        d.save(str(docx_p))
        r_docx = read_any_file(docx_p)
        check("readers: docx", r_docx["ok"] and "DOCX test" in r_docx["text"], f"kind={r_docx['kind']}")
    except Exception as exc:
        check("readers: docx", False, f"docx unavailable: {exc}")

    # xlsx
    xlsx_p = fixtures / "sample.xlsx"
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["client", "amount"])
        ws.append(["Ramesh", 25000])
        wb.save(str(xlsx_p))
        r_xlsx = read_any_file(xlsx_p)
        check("readers: xlsx", r_xlsx["ok"] and "Ramesh" in r_xlsx["text"], f"kind={r_xlsx['kind']}")
    except Exception as exc:
        check("readers: xlsx", False, f"openpyxl unavailable: {exc}")

    # zip listing
    zip_p = fixtures / "sample.zip"
    import zipfile

    with zipfile.ZipFile(zip_p, "w") as z:
        z.writestr("inside/hello.txt", "nested content")
    r_zip = read_any_file(zip_p)
    check("readers: archive listing", r_zip["ok"] and "inside/hello.txt" in r_zip["text"], f"kind={r_zip['kind']}")
    r_zip_member = read_any_file(zip_p, member="inside/hello.txt")
    check("readers: archive member", "nested content" in r_zip_member["text"], "member read ok")

    check("readers: supported list", "PDF" in supported_types() and "Excel" in supported_types(),
          supported_types()[:70])

    # fuzzy find + read by name (read_file tool)
    from jarvis.tools.file_tools import find_files, read_file

    uniq = fixtures / "nikhil_unique_report_9x.txt"
    uniq.write_text("line one\nunique content marker 42", encoding="utf-8")
    import os as _os
    _os.utime(uniq, None)

    f_out = await find_files.ainvoke({"name": "uniqe", "limit": 5})   # typo'd query
    r_exact = await read_file.ainvoke({"path": "nikhil_unique_report_9x"})
    r_typo = await read_file.ainvoke({"path": "nikhil_unqiue_reprot_9x"})  # typo'd, no path, no ext
    check("file_tools: find_files typo-tolerant", "nikhil_unique_report_9x" in f_out, f_out.splitlines()[0][:70])
    check("file_tools: read_file by fuzzy/typo name",
          "unique content marker 42" in r_exact and "unique content marker 42" in r_typo,
          "naam + typo dono se file padhi")


async def t_search_inside() -> None:
    from jarvis.tools.file_tools import search_inside_files

    assert FIXTURES_DIR is not None
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    (FIXTURES_DIR / "notes_one.txt").write_text("client Ramesh ka number 9812345678 hai", encoding="utf-8")
    out = await search_inside_files.ainvoke({"query": "Ramesh", "ext": "txt"})
    check("file_tools: content search (grep)", "notes_one.txt" in out and "9812345678" in out,
          out.splitlines()[0][:80])


def t_repl() -> None:
    from jarvis.runtime.repl import get_repl

    r = get_repl("selftest")
    r.reset()
    a = r.run("vals=[10,20,30]\nsum(vals)")
    check("repl: run + last expression", a["ok"] and "60" in a["stdout"], a["stdout"].strip()[:40])

    b = r.run("vals.append(40)\nsum(vals)")
    check("repl: variables persist across calls", b["ok"] and "100" in b["stdout"], b["stdout"].strip()[:40])

    c = r.run("1/0")
    check("repl: error returns traceback", (not c["ok"]) and "ZeroDivisionError" in c["stderr"], "ZeroDivisionError")

    d = r.run("import shutil; shutil.rmtree('/')")
    check("repl: dangerous code blocked", (not d["ok"]) and "blocked" in d["stderr"], d["stderr"][:60])

    t0 = time.time()
    e = r.run("while True: pass", timeout=2)
    check("repl: timeout returns fast (no hang)", (time.time() - t0) < 6 and not e["ok"],
          f"{time.time() - t0:.1f}s")


def t_session_state() -> None:
    from jarvis.memory.session_state import SessionState

    assert WORKSPACE_DIR is not None
    st = SessionState()
    st.set_workspace(WORKSPACE_DIR)
    st.add_task("selftest task", workers=["code"], status="running")
    st.remember_error("sample error for selftest")
    st.set_pending_steps(["step1", "step2"])
    st.note_app_opened("notepad")
    block = st.context_block()
    check("session_state: json persists context", "selftest task" in block and "sample error" in block,
          f"{len(block)} chars")
    check("session_state: file readable json", st.path.exists() and st.path.name.endswith("_memory.json"),
          str(st.path.name))
    sig = f"test::sig::{time.time_ns()}"
    a = st.seen_signature(sig)
    b = st.seen_signature(sig)
    check("session_state: loop signature counter", b == a + 1, f"repeat count {a} → {b} (loop detect ok)")


def t_process_manager() -> None:
    from jarvis.runtime.process_manager import list_processes, read_logs, start_process, stop_process

    inner = "import time; print(chr(112)+chr(114)+chr(111)+chr(99))"  # test fixture output
    cmd = f'"{sys.executable}" -c "{inner}"'
    info = start_process(cmd, name="selftest_proc")
    if info.popen is not None:
        info.popen.wait(timeout=5)
    procs = list_processes()
    logs = read_logs(info.pid, tail_lines=20)
    stopped = stop_process(info.pid)
    check("process_manager: start", any(p["pid"] == info.pid for p in procs), f"{len(procs)} procs")
    check("process_manager: logs capture", "proc" in logs or "alive" in logs, logs.strip().splitlines()[-1][:50])
    check("process_manager: stop", "band kar diya" in stopped or "exit" in stopped, stopped[:60])


async def t_approval_gate() -> None:
    from jarvis.security.gate import use_approval_handler
    from jarvis.tools.system_tools import run_shell_command

    old = os.environ.get("JARVIS_EXEC_APPROVAL")
    os.environ["JARVIS_EXEC_APPROVAL"] = "risky"
    try:
        with use_approval_handler(lambda _request: False):
            blocked = await run_shell_command.ainvoke({"command": "rm -rf ./some_random_dir_xyz"})
        check("approval: human denial prevents execution", "nahi chalaya" in blocked, blocked[:70])
        plain = await run_shell_command.ainvoke({"command": "echo jarvis-shell-ok"})
        check("approval: normal command runs without prompt", "jarvis-shell-ok" in plain,
              plain.strip().splitlines()[-1][:50])

        secret_env = {"JARVIS_SECRET_demo": "S3CR3T"}
        os.environ.update(secret_env)
        out = await run_shell_command.ainvoke({"command": "echo ${SECRET_demo}"})
        check("approval: secret expansion + redaction", "S3CR3T" not in out and "JARVIS_SECRET_demo" in out,
              out.strip().splitlines()[-1][:60])
    finally:
        if old is None:
            os.environ.pop("JARVIS_EXEC_APPROVAL", None)
        else:
            os.environ["JARVIS_EXEC_APPROVAL"] = old


def t_skills() -> None:
    from jarvis.skills.loader import catalog, matching_skills, skills_prompt_block

    cat = catalog()
    check("skills: catalog loads", "website_data_collection" in cat or "gst_invoice" in cat, cat.splitlines()[0][:70])
    m = matching_skills("browser pr alag alag website se price laao aur note me save kardo")
    check("skills: keyword matching", any(s.name == "website_data_collection" for s in m),
          f"matched={[s.name for s in m]}")
    block = skills_prompt_block("client invoice bana do")
    check("skills: prompt injection", "SKILL" in block or block == "", f"{len(block)} chars")


def t_router_local_vs_web() -> None:
    """'dhoondh' waala kaam local file ho to research me na jaye; web ho to research hi rahe."""
    from jarvis.dspy_engine.modules import fallback_route as fr

    check("router: local file task → desktop",
          fr("workspace me jo laptop wali notes file hai wo dhoondh ke padho") == ["desktop"],
          str(fr("workspace me jo laptop wali notes file hai wo dhoondh ke padho")))
    check("router: web task → research",
          "research" in fr("best laptop under 60000 dhoondho"),
          str(fr("best laptop under 60000 dhoondho")))
    check("router: web + note → dono workers",
          set(fr("web se laptop price laake note banao")) >= {"browser", "desktop"},
          str(fr("web se laptop price laake note banao")))
    check("router: weather/score → research",
          "research" in fr("delhi ka mausam batao") and "research" in fr("ipl ka latest score batao"),
          "mausam/score handled")
    check("router: app open → desktop",
          "desktop" in fr("safari khol do") and "desktop" in fr("vs code open karo"),
          "app open handled")


async def t_offline_file_pipeline() -> None:
    """Offline mode me bhi asli file dhoondh kar padhna (bina LLM) chalta ho."""
    from jarvis.graph.nodes import _extract_app_hint, _extract_file_hint, _local_file_find, _local_file_read

    check("offline: file hint ('laptop wali notes file')",
          _extract_file_hint("workspace me jo laptop wali notes file hai wo padho") == "laptop",
          _extract_file_hint("workspace me jo laptop wali notes file hai wo padho"))
    check("offline: file hint (extension)",
          _extract_file_hint("nikhil_report_9x.pdf padho bhai") == "nikhil_report_9x.pdf",
          _extract_file_hint("nikhil_report_9x.pdf padho bhai"))
    check("offline: app hint ('safari khol do')",
          _extract_app_hint("safari khol do") == "safari",
          _extract_app_hint("safari khol do"))

    # fixture khud banao in the temporary workspace; no repository workspace access.
    from jarvis.config import NOTES_DIR as configured_notes_dir

    notes_dir = configured_notes_dir
    notes_dir.mkdir(parents=True, exist_ok=True)
    suffix = str(time.time_ns())
    fixture_stem = f"laptop_prices_{suffix}"
    fixture = notes_dir / f"{fixture_stem}.md"
    from jarvis.tools.file_tools import invalidate_index

    fixture.write_text("# Laptop prices\n\nDell i5 ₹54,990 | HP Ryzen ₹49,990\n", encoding="utf-8")
    invalidate_index()
    try:
        typo_query = f"laptop_pricse_{suffix}"
        hits = await _local_file_find(typo_query)         # do typos ek saath
        check("offline: typo file find (laptop_pricse → laptop_prices)",
              bool(hits) and fixture.name in hits[0]["name"],
              f"{hits[0]['name']} @ {hits[0]['score']}%" if hits else "no hit")
        body = await _local_file_read(hits[0]["path"]) if hits else ""
        check("offline: file content actually read (asli data ke saath)",
              "54,990" in body, f"{len(body)} chars, content ok" if "54,990" in body else body[:80])
    finally:
        fixture.unlink(missing_ok=True)
        invalidate_index()


async def t_regressions_notes_cli() -> None:
    """Is round ke real bugs ke regression tests (note content leak, CLI markup crash, stale index)."""
    from jarvis.graph.nodes import _extract_note_body, _extract_note_name

    # 1) note content me instruction leak nahi hona chahiye
    body = _extract_note_body("notes me smoke_test naam se ye save kardo: CLI smoke verify")
    check("regression: note body me instruction leak nahi", body == "CLI smoke verify", repr(body))
    body2 = _extract_note_body("notes me laptop_deals naam se ye save kardo — Ram 16GB wala 55k")
    check("regression: note body em-dash separator", body2 == "Ram 16GB wala 55k", repr(body2))
    check("regression: note naam 'todo naam rakh ke' se sahi",
          _extract_note_name("jarvis, notes me todo naam rakh ke save kardo doodh lana") == "todo",
          _extract_note_name("jarvis, notes me todo naam rakh ke save kardo doodh lana"))

    # 2) CLI safe-print: file path jaisa `[/x/y]` bracket Rich markup crash na kare
    import io
    from contextlib import redirect_stdout

    from jarvis.cli import ui

    buf = io.StringIO()
    try:
        with redirect_stdout(buf):
            ui.cprint_safe("[/home/user/repo/skills/gst_invoice/SKILL.md]")
        crashed = False
    except Exception as exc:  # noqa: BLE001
        crashed = True
        err = str(exc)
    check("regression: CLI cprint_safe bracket path par crash nahi",
          not crashed, "" if not crashed else err[:80])

    # speech: emoji ka mapped word duplicate na ho (TTS me repeat sunai deta tha)
    from jarvis.voice.speech import for_speech

    spoken = for_speech("Ho gaya boss ✅ — Nikhil Mehra ke liye kaam complete.")
    check("regression: TTS me 'ho gaya' repeat nahi", spoken.lower().count("ho gaya") == 1, spoken[:60])
    check("regression: TTS me emoji/markdown nahi", "✅" not in spoken and "**" not in spoken, spoken[:60])

    # 3) Jarvis khud file likhe → turant dhoondh bhi le (index invalidate hota hai)
    from jarvis.config import NOTES_DIR as configured_notes_dir

    notes_dir = configured_notes_dir
    notes_dir.mkdir(parents=True, exist_ok=True)
    fresh_name = f"zz_regression_fresh_{time.time_ns()}"
    fresh = notes_dir / f"{fresh_name}.md"
    try:
        from jarvis.graph.nodes import _local_file_find
        from jarvis.tools.file_tools import invalidate_index
        from jarvis.tools.notes_tools import save_note

        await _local_file_find("laptop")          # pehle index warm karo (cache ban jaye)
        await save_note.ainvoke({"name": fresh_name, "content": "fresh content 12345"})
        hits = await _local_file_find(fresh_name)
        check("regression: apne save_note ki file turant mile (index invalidate)",
              bool(hits) and hits[0]["name"] == fresh.name,
              f"{hits[0]['name']} @ {hits[0]['score']}%" if hits else "no hit")
    finally:
        fresh.unlink(missing_ok=True)
        from jarvis.tools.file_tools import invalidate_index

        invalidate_index()


async def t_time_intent() -> None:
    """Spec: date/time kabhi hallucinate nahi — local clock fast-path se aaye."""
    from jarvis.runtime.agent import _is_time_question, _time_answer

    check("time: 'aaj ka date' detect", _is_time_question("aaj ka date aur time batao"))
    check("time: 'abhi kitne baje' detect", _is_time_question("abhi kitne baje hain"))
    check("time: weather hijack NAHI", not _is_time_question("aaj ka mausam batao"))
    check("time: browser task hijack NAHI", not _is_time_question("browser se laptop prices laao"))
    ans = _time_answer("aaj ka date aur time batao")
    import datetime as _dt
    from zoneinfo import ZoneInfo

    from jarvis.config import TIMEZONE

    today = _dt.datetime.now(ZoneInfo(TIMEZONE)).strftime("%d %B %Y")
    check("time: answer local clock se (aaj ki date match)", today in ans, ans.splitlines()[0][:60])
    check("time: timezone included only for requested time", "IST" in ans, ans.splitlines()[0][:60])
    kal = _time_answer("kal ki date kya hai")
    tomorrow = (_dt.datetime.now(ZoneInfo(TIMEZONE)) + _dt.timedelta(days=1)).strftime("%d %B %Y")
    check("time: kal ki date sahi", tomorrow in kal, kal.splitlines()[0][:60])

    from jarvis.graph.state import evidence_digest, local_clock_line

    check("time: routine evidence has no LOCAL CLOCK", "LOCAL CLOCK" not in evidence_digest({"findings": []}),
          local_clock_line()[:55])
    check("time: explicit schedule evidence includes local clock",
          "LOCAL CLOCK" in evidence_digest({"request": "Remind me tomorrow", "findings": []}),
          local_clock_line()[:55])

    # graph routing: pure time sawaal par web research nahi
    from jarvis.dspy_engine.modules import fallback_route

    check("time: route research/browser nahi karta", fallback_route("aaj ka date batao") == ["answer"],
          str(fallback_route("aaj ka date batao")))


def t_app_match_scores() -> None:
    """App search scoring — galat app kholna sabse bada bug tha (alias 100% inflation)."""
    from jarvis.tools.app_finder import _similarity

    def app(name: str, extra: str = "") -> dict:
        return {"name": name, "norm": name.lower(), "path": "desktop:" + name.split()[0].lower(),
                "kind": "desktop", "extra": extra}

    np_, im, ch = app("Notepad"), app("ImageMagick", "image;view;"), app("Google Chrome", "browser;")
    check("app-score: 'notepad' vs Notepad = 100", round(_similarity("notepad", np_)) == 100,
          f"{_similarity('notepad', np_):.0f}%")
    # REGRESSION: alias expansion har app ko 'notepad' se match kara deta tha (100%)
    check("app-score: 'notepad' vs ImageMagick LOW (alias bug fix)",
          _similarity("notepad", im) < 60, f"{_similarity('notepad', im):.0f}%")
    check("app-score: 'chrome' vs Google Chrome = 90", round(_similarity("chrome", ch)) == 90,
          f"{_similarity('chrome', ch):.0f}%")
    check("app-score: 'chrome' vs ImageMagick LOW", _similarity("chrome", im) < 60,
          f"{_similarity('chrome', im):.0f}%")
    check("app-score: 'gimp' vs Notepad LOW", _similarity("gimp", np_) < 60, f"{_similarity('gimp', np_):.0f}%")
    check("app-score: 'notep' (adhoora) vs Notepad ≥ 88", _similarity("notep", np_) >= 88,
          f"{_similarity('notep', np_):.0f}%")
    check("app-score: alias 'browser' → Chrome ≥ 88", _similarity("browser", ch) >= 88,
          f"{_similarity('browser', ch):.0f}%")
    check("app-score: alias 'notepad' → Text Editor ≥ 88 (indirect match)",
          _similarity("notepad", app("Text Editor")) >= 88, f"{_similarity('notepad', app('Text Editor')):.0f}%")
    check("app-score: 'vscode' → Visual Studio Code ≥ 88",
          _similarity("vscode", app("Visual Studio Code")) >= 88, "alias")


def t_offline_helpers() -> None:
    """E2E me mile teen bugs ke regression tests."""
    from jarvis.graph.nodes import _ext_filter, _extract_note_name
    from jarvis.tools.app_finder import _ratio, _similarity

    check("helper: note naam 'naam X se' sahi nikle (karo.md bug)",
          _extract_note_name("notepad me note save karo naam laptop_deals aur content likho X") == "laptop_deals",
          _extract_note_name("notepad me note save karo naam laptop_deals"))
    check("helper: note naam bina diye default",
          _extract_note_name("ye note save karo") == "jarvis_note", "default ok")
    check("helper: type filter ('resume pdf' → .pdf)",
          _ext_filter("meri resume pdf dhoondho") == (".pdf",), str(_ext_filter("meri resume pdf dhoondho")))
    check("helper: app ratio scale 0..1 (2428% bug)",
          _ratio("safari", "safari") == 1.0 and _ratio("abc", "xyz") == 0.0,
          f"self={_ratio('safari', 'safari')} other={_ratio('abc', 'xyz')}")
    check("helper: weak app match < 60 (galat app launch na ho)",
          _similarity("safari", {"name": "R", "norm": "r", "extra": "", "path": "/usr/bin/R"}) < 60,
          str(round(_similarity("safari", {"name": "R", "norm": "r", "extra": "", "path": "/usr/bin/R"}), 1)))


async def main() -> int:
    global FIXTURES_DIR, MEMORY_DB, NOTES_DIR, WORKSPACE_DIR

    _results.clear()
    _WEB_SEARCH_CALLS.clear()
    _WEB_FETCH_CALLS.clear()
    print("\n" + "=" * 66)
    print("  NikhilAI Jarvis — OFFLINE SELFTEST (temporary workspace; no network)")
    print("=" * 66)

    temp_dir = tempfile.TemporaryDirectory(prefix="nikhilai-selftest-", dir=str(_temporary_parent()))
    temp_root = Path(temp_dir.name).resolve()
    temp_outside_repo = not _is_within(temp_root, ROOT)
    previous_environment: dict[str, str | None] = {}
    network_attempts: list[str] = []
    restore_network = None
    paths_isolated = False
    database_created = False

    try:
        previous_environment = _configure_selftest_environment(temp_root)
        network_attempts, restore_network = _install_network_guard()

        # Import config only after all data paths are redirected; settings creates its dirs at import.
        from contextlib import ExitStack
        from unittest.mock import patch

        from jarvis.config import MEMORY_DB as configured_db
        from jarvis.config import NOTES_DIR as configured_notes
        from jarvis.config import WORKSPACE_DIR as configured_workspace
        from jarvis.tools import browser_tools, web_tools

        WORKSPACE_DIR = configured_workspace
        NOTES_DIR = configured_notes
        MEMORY_DB = configured_db
        FIXTURES_DIR = configured_workspace / "artifacts" / "selftest_fixtures"
        paths_isolated = temp_outside_repo and all(
            _is_within(path, temp_root)
            for path in (configured_workspace, configured_notes, configured_db)
        )

        if paths_isolated:
            async def no_live_browser_agent(*_args, **_kwargs):
                return None

            def mock_weather_tool(city: str = "") -> str:
                return _mock_weather(city)

            with ExitStack() as mocks:
                mocks.enter_context(patch.object(web_tools, "web_search_results", _mock_web_search))
                mocks.enter_context(patch.object(web_tools, "fetch_url_text", _mock_web_fetch))
                mocks.enter_context(patch.object(web_tools, "_weather_sync", mock_weather_tool))
                mocks.enter_context(patch.object(browser_tools, "build_browser_agent", no_live_browser_agent))

                print("\n[1] memory")
                await t_memory()
                print("\n[2] notes / notepad")
                await t_notes()
                print("\n[3] code engine (self-write + run + heal)")
                await t_code_engine()
                print("\n[4] browser-use integration")
                await t_browser()
                print("\n[5] DSPy reasoning modules")
                await t_dspy()
                print("\n[6] LangGraph supervisor (end-to-end)")
                await t_graph()
                print("\n[7] identity migration")
                t_identity()
                print("\n[8] app finder (kisi bhi app ko naam se kholna)")
                t_app_finder()
                print("\n[9] any-file reader (PDF/DOCX/XLSX/CSV/JSON/ZIP)")
                await t_file_readers()
                print("\n[10] content search (files ke andar)")
                await t_search_inside()
                print("\n[11] persistent Python REPL (self-correction sandbox)")
                t_repl()
                print("\n[12] short-term state (jarvis_memory.json)")
                t_session_state()
                print("\n[13] background process manager")
                t_process_manager()
                print("\n[14] exec approval + secrets")
                await t_approval_gate()
                print("\n[15] skills (SKILL.md workflows)")
                t_skills()
                print("\n[16] router — local file vs web intent")
                t_router_local_vs_web()
                print("\n[17] offline file pipeline (bina LLM bhi kaam)")
                await t_offline_file_pipeline()
                await t_regressions_notes_cli()
                await t_time_intent()
                t_app_match_scores()
                print("\n[18] offline helpers (regression)")
                t_offline_helpers()
        else:
            print("\nRefusing to run: configured workspace, notes, or database escaped the temp directory.")
    finally:
        if MEMORY_DB is not None:
            database_created = MEMORY_DB.is_file()
        if restore_network is not None:
            restore_network()
        _restore_environment(previous_environment)
        temp_dir.cleanup()

    temp_removed = not temp_root.exists()
    no_network = not network_attempts
    hermetic = paths_isolated and database_created and temp_removed and no_network
    check(
        "selftest: temp workspace/DB cleaned; repository path unused; no live network",
        hermetic,
        f"isolated={paths_isolated} temp_outside_repo={temp_outside_repo} db_created={database_created} "
        f"temp_removed={temp_removed} network_attempts={len(network_attempts)}",
    )

    passed = sum(1 for _n, ok, _d in _results if ok)
    total = len(_results)
    print("\n" + "=" * 66)
    print(f"  RESULT: {passed}/{total} checks passed  ({'ALL GOOD 🚀' if passed == total else 'kuch fail hue ⚠️'})")
    print("=" * 66 + "\n")
    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
